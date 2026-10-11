from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook

from app.agents.document_ingestion import DocumentIngestionError, extract_document


class Upload(BytesIO):
    def __init__(self, name: str, content: bytes):
        super().__init__(content)
        self.name = name


def _text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = b"%PDF-1.4\n"
    offsets = []
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(output)
    output += b"xref\n0 6\n0000000000 65535 f \n"
    output += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    output += (
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref).encode()
        + b"\n%%EOF"
    )
    return output


def test_txt_preserves_metadata_and_section() -> None:
    content = b"Annual family income must not exceed Rs. 250000."
    result = extract_document(Upload("policy.txt", content))

    assert result.filename == "policy.txt"
    assert result.file_size == len(content)
    assert result.sha256
    assert result.document_id.startswith("document-")
    assert result.sections[0]["section"] == 1
    assert result.sections[0]["text"] == content.decode()


def test_docx_has_section_location_but_no_fabricated_page() -> None:
    document = Document()
    document.add_paragraph("Minimum attendance of 75% is required.")
    buffer = BytesIO()
    document.save(buffer)

    result = extract_document(Upload("rules.docx", buffer.getvalue()))

    assert result.sections[0]["section"] == 1
    assert "page" not in result.sections[0]


def test_pdf_preserves_real_page_number() -> None:
    result = extract_document(
        Upload("policy.pdf", _text_pdf("Annual family income must not exceed Rs. 250000."))
    )

    assert result.sections[0]["page"] == 1
    assert "250000" in result.sections[0]["text"]


@pytest.mark.parametrize(
    ("name", "content", "code"),
    [
        ("notes.csv", b"value", "unsupported_type"),
        ("empty.txt", b"", "empty_file"),
        ("broken.pdf", b"not a PDF", "corrupt_document"),
    ],
)
def test_invalid_uploads_are_reported(name: str, content: bytes, code: str) -> None:
    with pytest.raises(DocumentIngestionError) as error:
        extract_document(Upload(name, content))

    assert error.value.code == code


def test_scanned_pdf_is_reported_without_claiming_ocr() -> None:
    from pypdf import PdfWriter

    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)

    with pytest.raises(DocumentIngestionError) as error:
        extract_document(Upload("scan.pdf", buffer.getvalue()))

    assert error.value.code == "no_extractable_text"
    assert "OCR is not implemented" in error.value.message


def test_size_limit_is_enforced() -> None:
    with pytest.raises(DocumentIngestionError) as error:
        extract_document(
            Upload("large.txt", b"12345"),
            max_file_size_bytes=4,
        )

    assert error.value.code == "file_too_large"


def _workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Eligibility"
    first.append(["Requirement", "Value"])
    first.append(["Annual family income must not exceed", "₹2,50,000"])
    first.append([])
    first.append(["Minimum attendance", "75%"])
    second = workbook.create_sheet("Documents")
    second.append(["Document", "Required"])
    second.append(["Income certificate", True])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_xlsx_extracts_multiple_worksheets_headers_rows_and_coordinates() -> None:
    content = _workbook_bytes()
    result = extract_document(Upload("policy.xlsx", content))

    assert result.file_type == "xlsx"
    assert {section["worksheet_name"] for section in result.sections} == {
        "Eligibility",
        "Documents",
    }
    eligibility_rows = [
        section
        for section in result.sections
        if section["worksheet_name"] == "Eligibility"
    ]
    assert [section["row_number"] for section in eligibility_rows] == [1, 2, 4]
    assert eligibility_rows[1]["headers"] == ["Requirement", "Value"]
    assert eligibility_rows[1]["cells"] == [
        {"coordinate": "A2", "column": "Requirement", "value": "Annual family income must not exceed"},
        {"coordinate": "B2", "column": "Value", "value": "₹2,50,000"},
    ]
    assert "Requirement: Annual family income must not exceed" in eligibility_rows[1]["text"]


def test_xlsx_empty_worksheet_is_preserved_without_fabricated_text() -> None:
    workbook = Workbook()
    workbook.active.title = "Empty"
    buffer = BytesIO()
    workbook.save(buffer)

    result = extract_document(Upload("empty.xlsx", buffer.getvalue()))

    assert len(result.sections) == 1
    assert result.sections[0]["worksheet_name"] == "Empty"
    assert result.sections[0]["worksheet_empty"] is True
    assert result.sections[0]["text"] == ""


def test_corrupt_xlsx_is_reported() -> None:
    with pytest.raises(DocumentIngestionError) as error:
        extract_document(Upload("broken.xlsx", b"not an xlsx"))

    assert error.value.code == "corrupt_document"
