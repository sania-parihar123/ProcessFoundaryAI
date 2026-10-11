"""Local, provenance-preserving extraction for supported document uploads."""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

DEFAULT_MAX_FILE_SIZE_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".xlsx"})


class DocumentIngestionError(ValueError):
    """A user-facing validation or extraction error for one upload."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ExtractedDocument:
    """Extracted text and metadata for one successfully read document."""

    document_id: str
    filename: str
    file_type: str
    file_size: int
    sha256: str
    sections: list[dict[str, Any]]

    def metadata(self) -> dict[str, Any]:
        return {
            "document_id": self.document_id,
            "filename": self.filename,
            "file_type": self.file_type,
            "file_size": self.file_size,
            "sha256": self.sha256,
        }


def _upload_name(upload: Any) -> str:
    name = getattr(upload, "name", None)
    if not isinstance(name, str) or not name.strip():
        raise DocumentIngestionError(
            "invalid_filename", "Each upload must have a filename."
        )
    return name.strip()


def _upload_bytes(upload: Any) -> bytes:
    if isinstance(upload, (bytes, bytearray)):
        return bytes(upload)
    getter = getattr(upload, "getvalue", None)
    if callable(getter):
        value = getter()
    else:
        reader = getattr(upload, "read", None)
        if not callable(reader):
            raise DocumentIngestionError(
                "invalid_upload", "Upload does not provide readable bytes."
            )
        value = reader()
    if not isinstance(value, (bytes, bytearray)):
        raise DocumentIngestionError(
            "invalid_upload", "Upload did not provide binary content."
        )
    return bytes(value)


def _section(
    *,
    document_id: str,
    filename: str,
    text: str,
    location: dict[str, int],
) -> dict[str, Any]:
    location_name, location_value = next(iter(location.items()))
    return {
        "section_id": f"{document_id}:{location_name}-{location_value}",
        "document_id": document_id,
        "filename": filename,
        location_name: location_value,
        "text": text,
    }


def _extract_txt(document_id: str, filename: str, content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentIngestionError(
            "invalid_text_encoding",
            "TXT files must be UTF-8 encoded.",
        ) from exc
    if not text.strip():
        return []
    return [_section(document_id=document_id, filename=filename, text=text, location={"section": 1})]


def _extract_docx(document_id: str, filename: str, content: bytes) -> list[dict[str, Any]]:
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise DocumentIngestionError(
            "corrupt_document", "The DOCX file could not be read."
        ) from exc
    sections: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if text:
            sections.append(
                _section(
                    document_id=document_id,
                    filename=filename,
                    text=text,
                    location={"section": index},
                )
            )
    return sections


def _extract_pdf(document_id: str, filename: str, content: bytes) -> list[dict[str, Any]]:
    try:
        reader = PdfReader(BytesIO(content))
        sections = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                sections.append(
                    _section(
                        document_id=document_id,
                        filename=filename,
                        text=text,
                        location={"page": page_number},
                    )
                )
    except Exception as exc:
        raise DocumentIngestionError(
            "corrupt_document", "The PDF file could not be read."
        ) from exc
    if not sections:
        raise DocumentIngestionError(
            "no_extractable_text",
            "The PDF contains no extractable text. OCR is not implemented.",
        )
    return sections


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip()


def _worksheet_section(
    *,
    document_id: str,
    filename: str,
    worksheet_name: str,
    row_number: int | None,
    headers: list[str],
    cells: list[dict[str, Any]],
    text: str,
    empty: bool = False,
) -> dict[str, Any]:
    row_part = f"row-{row_number}" if row_number is not None else "empty"
    return {
        "section_id": f"{document_id}:worksheet-{worksheet_name}:{row_part}",
        "document_id": document_id,
        "filename": filename,
        "worksheet_name": worksheet_name,
        "row_number": row_number,
        "headers": headers,
        "cells": cells,
        "text": text,
        "worksheet_empty": empty,
    }


def _extract_xlsx(
    document_id: str,
    filename: str,
    content: bytes,
) -> list[dict[str, Any]]:
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise DocumentIngestionError(
            "corrupt_document", "The XLSX workbook could not be read."
        ) from exc

    sections: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets:
            populated_rows: list[tuple[int, list[Any]]] = []
            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                values = [cell.value for cell in row]
                if any(_cell_text(value) for value in values):
                    populated_rows.append((row_number, values))

            if not populated_rows:
                sections.append(
                    _worksheet_section(
                        document_id=document_id,
                        filename=filename,
                        worksheet_name=worksheet.title,
                        row_number=None,
                        headers=[],
                        cells=[],
                        text="",
                        empty=True,
                    )
                )
                continue

            _, header_values = populated_rows[0]
            headers = [
                _cell_text(value) or f"Column {index}"
                for index, value in enumerate(header_values, start=1)
            ]
            for row_number, values in populated_rows:
                cells = []
                pairs = []
                for index, value in enumerate(values, start=1):
                    text_value = _cell_text(value)
                    if not text_value:
                        continue
                    coordinate = f"{worksheet.cell(row_number, index).column_letter}{row_number}"
                    header = headers[index - 1]
                    cells.append(
                        {
                            "coordinate": coordinate,
                            "column": header,
                            "value": value,
                        }
                    )
                    pairs.append(f"{header}: {text_value}")
                sections.append(
                    _worksheet_section(
                        document_id=document_id,
                        filename=filename,
                        worksheet_name=worksheet.title,
                        row_number=row_number,
                        headers=headers,
                        cells=cells,
                        text="; ".join(pairs),
                    )
                )
    finally:
        workbook.close()
    return sections


def extract_document(
    upload: Any,
    *,
    max_file_size_bytes: int = DEFAULT_MAX_FILE_SIZE_BYTES,
) -> ExtractedDocument:
    """Validate and extract one PDF, DOCX, XLSX, or UTF-8 TXT upload."""

    filename = _upload_name(upload)
    suffix = Path(filename).suffix.casefold()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentIngestionError(
            "unsupported_type",
            f"Unsupported file type '{suffix or 'unknown'}'. Use PDF, DOCX, XLSX, or TXT.",
        )
    if max_file_size_bytes <= 0:
        raise ValueError("max_file_size_bytes must be positive.")
    content = _upload_bytes(upload)
    file_size = len(content)
    if file_size == 0:
        raise DocumentIngestionError("empty_file", "The uploaded file is empty.")
    if file_size > max_file_size_bytes:
        raise DocumentIngestionError(
            "file_too_large",
            f"File exceeds the {max_file_size_bytes} byte size limit.",
        )

    sha256 = hashlib.sha256(content).hexdigest()
    document_id = f"document-{sha256[:16]}"
    if suffix == ".txt":
        sections = _extract_txt(document_id, filename, content)
    elif suffix == ".docx":
        sections = _extract_docx(document_id, filename, content)
    elif suffix == ".xlsx":
        sections = _extract_xlsx(document_id, filename, content)
    else:
        sections = _extract_pdf(document_id, filename, content)
    if not sections:
        raise DocumentIngestionError(
            "no_extractable_text",
            "The document contains no extractable text.",
        )
    return ExtractedDocument(
        document_id=document_id,
        filename=filename,
        file_type=suffix[1:],
        file_size=file_size,
        sha256=sha256,
        sections=sections,
    )
